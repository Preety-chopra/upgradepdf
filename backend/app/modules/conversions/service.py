from __future__ import annotations

import csv
import shutil
import subprocess
import zipfile
from pathlib import Path
from typing import Iterable, Literal

import fitz  # PyMuPDF
import pandas as pd
import pdfplumber
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from fastapi import HTTPException
from PIL import Image, ImageOps

Image.MAX_IMAGE_PIXELS = 120_000_000

PdfToWordMode = Literal["auto", "layout", "text", "image", "ocr"]


def _zip_files(zip_path: Path, files: Iterable[Path]) -> Path:
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for file_path in files:
            zf.write(file_path, arcname=file_path.name)
    return zip_path


def pdf_to_jpg_zip(pdf_path: Path, output_dir: Path, *, dpi: int = 200, quality: int = 90) -> Path:
    if dpi < 72 or dpi > 400:
        raise HTTPException(status_code=400, detail="DPI must be between 72 and 400.")
    if quality < 50 or quality > 100:
        raise HTTPException(status_code=400, detail="JPEG quality must be between 50 and 100.")

    image_dir = output_dir / "jpg_pages"
    image_dir.mkdir(parents=True, exist_ok=True)
    jpg_files: list[Path] = []

    try:
        with fitz.open(pdf_path) as doc:
            if doc.page_count == 0:
                raise HTTPException(status_code=400, detail="PDF has no pages.")

            zoom = dpi / 72
            matrix = fitz.Matrix(zoom, zoom)
            for page_index, page in enumerate(doc, start=1):
                pix = page.get_pixmap(matrix=matrix, alpha=False)
                image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                jpg_path = image_dir / f"{pdf_path.stem}_page_{page_index:03d}.jpg"
                image.save(jpg_path, "JPEG", quality=quality, optimize=True)
                jpg_files.append(jpg_path)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Unable to render PDF pages: {exc}") from exc

    return _zip_files(output_dir / f"{pdf_path.stem}_jpg_pages.zip", jpg_files)


def images_to_pdf(image_paths: list[Path], output_pdf: Path) -> Path:
    if not image_paths:
        raise HTTPException(status_code=400, detail="Please upload at least one image.")

    pages: list[Image.Image] = []
    try:
        for image_path in image_paths:
            with Image.open(image_path) as img:
                img = ImageOps.exif_transpose(img)
                if img.mode in ("RGBA", "LA"):
                    background = Image.new("RGB", img.size, "white")
                    background.paste(img, mask=img.getchannel("A"))
                    img = background
                else:
                    img = img.convert("RGB")
                pages.append(img.copy())

        first, *rest = pages
        output_pdf.parent.mkdir(parents=True, exist_ok=True)
        first.save(output_pdf, "PDF", save_all=True, append_images=rest, resolution=100.0)
        return output_pdf
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Unable to convert images to PDF: {exc}") from exc
    finally:
        for page in pages:
            try:
                page.close()
            except Exception:
                pass


def office_to_pdf(input_path: Path, output_dir: Path) -> Path:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise HTTPException(
            status_code=500,
            detail="LibreOffice is not installed in the API container. Install it in Dockerfile for Word/Excel to PDF.",
        )

    output_dir.mkdir(parents=True, exist_ok=True)
    profile_dir = output_dir / "lo-profile"
    profile_dir.mkdir(parents=True, exist_ok=True)

    command = [
        soffice,
        "--headless",
        "--nologo",
        "--nofirststartwizard",
        "--norestore",
        f"-env:UserInstallation={profile_dir.as_uri()}",
        "--convert-to",
        "pdf",
        "--outdir",
        str(output_dir),
        str(input_path),
    ]

    result = subprocess.run(command, capture_output=True, text=True, timeout=120)
    if result.returncode != 0:
        raise HTTPException(
            status_code=500,
            detail=(result.stderr or result.stdout or "LibreOffice conversion failed.").strip(),
        )

    converted = output_dir / f"{input_path.stem}.pdf"
    if not converted.exists():
        candidates = sorted(output_dir.glob("*.pdf"))
        if not candidates:
            raise HTTPException(status_code=500, detail="LibreOffice did not produce a PDF output file.")
        converted = candidates[0]

    return converted


def _pdf_text_stats(pdf_path: Path) -> dict[str, int | float]:
    """Return quick text-layer stats used to decide whether the PDF is scanned."""
    with fitz.open(pdf_path) as pdf:
        if pdf.page_count == 0:
            raise HTTPException(status_code=400, detail="PDF has no pages.")

        page_count = pdf.page_count
        text_pages = 0
        text_chars = 0
        image_pages = 0
        for page in pdf:
            text = page.get_text("text").strip()
            if text:
                text_pages += 1
                text_chars += len(text)
            if page.get_images(full=True):
                image_pages += 1

    return {
        "page_count": page_count,
        "text_pages": text_pages,
        "text_chars": text_chars,
        "image_pages": image_pages,
        "text_page_ratio": text_pages / page_count if page_count else 0,
    }


def _looks_like_scanned_pdf(stats: dict[str, int | float]) -> bool:
    """
    A scanned/image PDF usually has very few extractable text pages compared with total pages.
    Mixed documents are treated as scanned if the text layer is only present on a tiny minority of pages.
    """
    page_count = int(stats["page_count"])
    text_pages = int(stats["text_pages"])
    text_chars = int(stats["text_chars"])
    text_page_ratio = float(stats["text_page_ratio"])

    if page_count == 0:
        return False
    if text_pages == 0:
        return True
    if text_page_ratio < 0.25 and text_chars < max(1200, page_count * 80):
        return True
    return False


def _page_margin_pt() -> float:
    # Small margin keeps full-page screenshots inside Word's printable area.
    return 12.0


def _configure_section_for_pdf_page(section, page_rect) -> None:
    margin = _page_margin_pt()
    section.page_width = Pt(float(page_rect.width))
    section.page_height = Pt(float(page_rect.height))
    section.top_margin = Pt(margin)
    section.bottom_margin = Pt(margin)
    section.left_margin = Pt(margin)
    section.right_margin = Pt(margin)
    section.header_distance = Pt(0)
    section.footer_distance = Pt(0)


def _render_page_png(page, output_dir: Path, page_index: int, *, dpi: int) -> Path:
    if dpi < 96 or dpi > 300:
        raise HTTPException(status_code=400, detail="PDF-to-Word image/OCR DPI must be between 96 and 300.")

    output_dir.mkdir(parents=True, exist_ok=True)
    zoom = dpi / 72
    pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
    image_path = output_dir / f"page_{page_index:04d}.png"
    pix.save(str(image_path))
    return image_path


def pdf_to_image_docx(pdf_path: Path, output_docx: Path, work_dir: Path, *, dpi: int = 180) -> Path:
    """
    Convert each PDF page into one full-page image inside Word.

    This is the correct fallback for scanned PDFs. It preserves the visual page appearance,
    but the text will not be editable unless OCR is run separately.
    """
    output_docx.parent.mkdir(parents=True, exist_ok=True)
    image_dir = work_dir / "pdf_page_images"
    document = Document()

    try:
        with fitz.open(pdf_path) as pdf:
            if pdf.page_count == 0:
                raise HTTPException(status_code=400, detail="PDF has no pages.")

            for page_number, page in enumerate(pdf, start=1):
                if page_number == 1:
                    section = document.sections[0]
                else:
                    section = document.add_section(WD_SECTION.NEW_PAGE)

                _configure_section_for_pdf_page(section, page.rect)
                image_path = _render_page_png(page, image_dir, page_number, dpi=dpi)

                paragraph = document.paragraphs[-1] if page_number == 1 and document.paragraphs else document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1
                run = paragraph.add_run()

                available_width = max(float(page.rect.width) - (2 * _page_margin_pt()), 72.0)
                run.add_picture(str(image_path), width=Pt(available_width))

        document.save(output_docx)
        return output_docx
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Unable to create image-based Word file: {exc}") from exc


def pdf_to_text_docx(pdf_path: Path, output_docx: Path) -> Path:
    """Simple editable text extraction for PDFs that already contain a selectable text layer."""
    document = Document()
    document.add_heading("Extracted PDF Text", level=1)
    document.add_paragraph(f"Source file: {pdf_path.name}")

    try:
        with fitz.open(pdf_path) as pdf:
            if pdf.page_count == 0:
                raise HTTPException(status_code=400, detail="PDF has no pages.")

            for page_index, page in enumerate(pdf, start=1):
                document.add_heading(f"Page {page_index}", level=2)
                text = page.get_text("text").strip()
                if text:
                    for block in text.split("\n\n"):
                        paragraph_text = "\n".join(line.rstrip() for line in block.splitlines()).strip()
                        if paragraph_text:
                            document.add_paragraph(paragraph_text)
                else:
                    document.add_paragraph("[No extractable text found on this page]")
                if page_index != pdf.page_count:
                    document.add_page_break()

        output_docx.parent.mkdir(parents=True, exist_ok=True)
        document.save(output_docx)
        return output_docx
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Unable to extract PDF text to Word: {exc}") from exc


def pdf_to_layout_docx(pdf_path: Path, output_docx: Path) -> Path:
    """
    Layout-aware PDF to DOCX for digital PDFs.

    pdf2docx works on PDFs that contain a real text/vector layer. It cannot recover editable
    text from scanned pages without OCR, so scanned PDFs should use image or OCR mode.
    """
    try:
        from pdf2docx import Converter
    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail="pdf2docx is not installed. Add pdf2docx to requirements.txt and rebuild the API container.",
        ) from exc

    output_docx.parent.mkdir(parents=True, exist_ok=True)
    converter = None
    try:
        converter = Converter(str(pdf_path))
        converter.convert(str(output_docx), start=0, end=None)
        return output_docx
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Unable to convert PDF layout to Word: {exc}") from exc
    finally:
        if converter is not None:
            try:
                converter.close()
            except Exception:
                pass


def pdf_to_ocr_docx(pdf_path: Path, output_docx: Path, work_dir: Path, *, dpi: int = 220) -> Path:
    """
    OCR scanned PDF pages into editable text.

    This produces editable text, but scanned tables/forms will be approximate. For visual fidelity,
    use image mode. For exact editable conversion, a commercial/OCR-layout engine is usually needed.
    """
    tesseract = shutil.which("tesseract")
    if not tesseract:
        raise HTTPException(
            status_code=500,
            detail="Tesseract OCR is not installed in the API container. Install tesseract-ocr in Dockerfile.",
        )

    try:
        import pytesseract
    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail="pytesseract is not installed. Add pytesseract to requirements.txt and rebuild the API container.",
        ) from exc

    image_dir = work_dir / "ocr_images"
    document = Document()
    document.add_heading("OCR Extracted PDF Text", level=1)
    document.add_paragraph(f"Source file: {pdf_path.name}")
    document.add_paragraph(
        "Note: This OCR output is editable text. Table layout and exact formatting may need manual correction."
    )

    try:
        with fitz.open(pdf_path) as pdf:
            if pdf.page_count == 0:
                raise HTTPException(status_code=400, detail="PDF has no pages.")

            for page_number, page in enumerate(pdf, start=1):
                document.add_heading(f"Page {page_number}", level=2)
                image_path = _render_page_png(page, image_dir, page_number, dpi=dpi)
                with Image.open(image_path) as image:
                    text = pytesseract.image_to_string(image, lang="eng", config="--psm 6").strip()
                    if not text:
                        text = pytesseract.image_to_string(image, lang="eng", config="--psm 11").strip()

                if text:
                    for block in text.split("\n\n"):
                        block = "\n".join(line.rstrip() for line in block.splitlines()).strip()
                        if block:
                            document.add_paragraph(block)
                else:
                    document.add_paragraph("[OCR could not detect readable text on this page]")

                if page_number != pdf.page_count:
                    document.add_page_break()

        output_docx.parent.mkdir(parents=True, exist_ok=True)
        document.save(output_docx)
        return output_docx
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Unable to OCR PDF to Word: {exc}") from exc


def pdf_to_docx(pdf_path: Path, output_docx: Path, work_dir: Path, *, mode: PdfToWordMode = "auto", dpi: int = 180) -> Path:
    """
    Main PDF to Word dispatcher.

    Modes:
    - auto: digital PDFs -> layout conversion; scanned PDFs -> visual page-image DOCX
    - layout: use pdf2docx for digital PDFs
    - text: extract selectable text only
    - image: preserve each page visually as an image in Word
    - ocr: OCR scanned pages into editable text
    """
    allowed_modes = {"auto", "layout", "text", "image", "ocr"}
    if mode not in allowed_modes:
        raise HTTPException(status_code=400, detail=f"Invalid PDF-to-Word mode. Use one of: {', '.join(sorted(allowed_modes))}.")

    stats = _pdf_text_stats(pdf_path)

    if mode == "text":
        return pdf_to_text_docx(pdf_path, output_docx)

    if mode == "image":
        return pdf_to_image_docx(pdf_path, output_docx, work_dir, dpi=dpi)

    if mode == "ocr":
        return pdf_to_ocr_docx(pdf_path, output_docx, work_dir, dpi=max(dpi, 200))

    if mode == "layout":
        return pdf_to_layout_docx(pdf_path, output_docx)

    # auto mode
    if _looks_like_scanned_pdf(stats):
        return pdf_to_image_docx(pdf_path, output_docx, work_dir, dpi=dpi)

    try:
        return pdf_to_layout_docx(pdf_path, output_docx)
    except HTTPException:
        # If layout conversion is unavailable or fails, still return a useful editable text DOCX.
        return pdf_to_text_docx(pdf_path, output_docx)


def _clean_table(raw_table: list[list[str | None]]) -> list[list[str]]:
    cleaned: list[list[str]] = []
    for row in raw_table:
        cleaned_row = [(cell or "").strip() if isinstance(cell, str) or cell is None else str(cell) for cell in row]
        if any(cell for cell in cleaned_row):
            cleaned.append(cleaned_row)
    return cleaned


def _make_dataframe(table: list[list[str]]) -> pd.DataFrame:
    if not table:
        return pd.DataFrame()

    width = max(len(row) for row in table)
    normalized = [row + [""] * (width - len(row)) for row in table]

    # Keep table extraction lossless. The first row may or may not be an actual header, so do not promote it.
    return pd.DataFrame(normalized)


def extract_pdf_tables(pdf_path: Path) -> list[tuple[str, pd.DataFrame]]:
    tables: list[tuple[str, pd.DataFrame]] = []
    try:
        with pdfplumber.open(pdf_path) as pdf:
            for page_number, page in enumerate(pdf.pages, start=1):
                for table_number, raw_table in enumerate(page.extract_tables() or [], start=1):
                    cleaned = _clean_table(raw_table)
                    df = _make_dataframe(cleaned)
                    if not df.empty:
                        tables.append((f"page_{page_number}_table_{table_number}", df))
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Unable to extract tables from PDF: {exc}") from exc

    if not tables:
        raise HTTPException(
            status_code=422,
            detail="No tables were detected in this PDF. Scanned PDFs need OCR/table detection before export.",
        )

    return tables


def pdf_to_xlsx(pdf_path: Path, output_xlsx: Path) -> Path:
    tables = extract_pdf_tables(pdf_path)
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output_xlsx, engine="openpyxl") as writer:
        for sheet_name, df in tables:
            safe_sheet = sheet_name[:31]
            df.to_excel(writer, sheet_name=safe_sheet, index=False, header=False)

    return output_xlsx


def pdf_to_csv_or_zip(pdf_path: Path, output_dir: Path) -> Path:
    tables = extract_pdf_tables(pdf_path)
    output_dir.mkdir(parents=True, exist_ok=True)

    csv_files: list[Path] = []
    for table_name, df in tables:
        csv_path = output_dir / f"{pdf_path.stem}_{table_name}.csv"
        df.to_csv(csv_path, index=False, header=False, quoting=csv.QUOTE_MINIMAL)
        csv_files.append(csv_path)

    if len(csv_files) == 1:
        return csv_files[0]

    return _zip_files(output_dir / f"{pdf_path.stem}_tables_csv.zip", csv_files)
